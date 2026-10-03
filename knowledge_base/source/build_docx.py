"""Build editable .docx originals mirroring the PDFs.
Run: python knowledge_base/source/build_docx.py
Requires: pip install python-docx
The .docx files are the editable source; PDFs are rebuilt via build_pdfs.py.
Page break is inserted exactly where [PAGE BREAK] is marked in the spec.
"""
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

FONT = "Calibri"


def style_doc(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    for hname, sz in [("Heading 1", 14), ("Heading 2", 11), ("Heading 3", 10.5)]:
        hs = doc.styles[hname]
        hs.font.name = FONT
        hs.font.size = Pt(sz)


def add_table(doc, header, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Light Grid Accent 1"
    for j, h in enumerate(header):
        c = t.cell(0, j)
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9)
        r.font.name = FONT
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            r = c.paragraphs[0].add_run(val)
            r.font.size = Pt(9)
            r.font.name = FONT


def build_sop(path):
    doc = Document()
    style_doc(doc)
    doc.add_heading("Aravalli Mobility Pvt Ltd", level=1)
    doc.add_heading("SOP-SC-014: Supply Disruption Response", level=2)
    doc.add_paragraph("Owner: Head of Supply Chain Management \u00b7 Version: 3.2 \u00b7 Applies to: Manesar plant, all direct-material categories.")
    doc.add_heading("Section 1 - Purpose and Scope", level=3)
    doc.add_paragraph("1.1 This SOP defines the mandatory steps when an external event (port congestion or closure, extreme weather, geopolitical tension, sanctions, supplier force majeure) threatens inbound supply of direct materials.")
    doc.add_paragraph("1.2 Every automated or manual action taken under this SOP must cite the clause that authorises it.")
    doc.add_heading("Section 2 - Severity Classification", level=3)
    add_table(doc,
              ["Level", "Definition", "Response window"],
              [["L1 - Watch", "Expected delay under 7 days, or event not yet confirmed by two independent sources.", "Monitor; re-assess within 24 h. No RFQ."],
               ["L2 - Material", "Expected delay of 7 to 21 days on a named port, lane or supplier.", "RFQ trigger test (Section 3) within 8 h."],
               ["L3 - Critical", "Expected delay over 21 days, port closure beyond 21 days, blockade, export ban, or supplier force majeure notice.", "RFQ trigger test within 4 h; contingency plan mandatory."]])
    doc.add_paragraph("2.2 Where the delay is uncertain, the analyst must use the upper end of the credible estimate (conservative planning principle).")
    doc.add_heading("Section 3 - RFQ Trigger Rule", level=3)
    doc.add_paragraph("3.1 For each affected part compute days of cover = on-hand quantity / daily consumption.")
    doc.add_paragraph("3.2 An RFQ to an approved backup supplier is mandatory when: days of cover < projected delay (days) + 5-day safety buffer.")
    doc.add_paragraph("3.3 Shortfall units = (projected delay + 5 - days of cover) x daily consumption. RFQ quantity = shortfall units x 1.2, rounded up to the nearest multiple of the backup supplier's MOQ.")
    doc.add_paragraph("3.4 RFQs may only be issued to suppliers listed as APPROVED_BACKUP for that part in the ERP and marked APPROVED on the Approved Vendor List (Vendor Contracts document, Schedule A). CONDITIONAL suppliers require Head of SCM sign-off before an RFQ is released.")
    doc.add_paragraph("3.5 For L3 events affecting a critical part, RFQs must go to at least two approved backups where two exist.")
    doc.add_paragraph("3.6 Parts whose days of cover exceed projected delay + 5 days are placed on L1 watch; no RFQ is raised.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Section 4 - Contingency Measures", level=3)
    doc.add_paragraph("4.1 Expedited freight: Air freight may be authorised for critical parts when the freight premium is below 15% of the affected PO value.")
    doc.add_paragraph("4.2 Safety stock uplift: During any active L2 or L3 event the safety stock target for affected parts is raised by 30% until the event is closed.")
    doc.add_paragraph("4.3 Re-routing: For port-specific events, the buyer shall request the primary supplier to re-route via an alternate export port (e.g. Kaohsiung to Keelung, or trans-shipment via Singapore) before cancelling open POs.")
    doc.add_paragraph("4.4 Open POs: Open POs with the disrupted supplier are not cancelled; they are placed on HOLD-REVIEW and re-dated.")
    doc.add_paragraph("4.5 Production protection: If projected stock-out falls within 7 days, Production Planning must be notified to re-sequence builds toward models not using the affected part.")
    doc.add_paragraph("4.6 Capacity shortfall: If a backup supplier's monthly capacity is below the RFQ quantity, the buyer must combine the RFQ with measures 4.1 or 4.3 for the balance.")
    doc.add_heading("Section 5 - Approval Matrix", level=3)
    doc.add_paragraph("5.1 Approval is based on the total estimated value of all RFQs raised in one action record.")
    add_table(doc,
              ["Total estimated RFQ value (INR)", "Approving authority"],
              [["Up to 25,00,000", "Operations Manager - Procurement"],
               ["25,00,001 to 1,00,00,000", "Head of Supply Chain Management"],
               ["Above 1,00,00,000", "Chief Financial Officer"]])
    doc.add_heading("Section 6 - Communication and Records", level=3)
    doc.add_paragraph("6.1 The Operations Coordinator must publish an action record (RFQ or contingency plan) to the procurement workflow system within the response window of Section 2.")
    doc.add_paragraph("6.2 The record must include: severity level, affected parts, days of cover, projected delay, shortfall, supplier chosen, quantity, approving authority and policy citations.")
    doc.add_paragraph("6.3 Records are retained for 7 years for audit.")
    doc.save(path)
    print("wrote", path)


def build_vendor(path):
    doc = Document()
    style_doc(doc)
    doc.add_heading("Aravalli Mobility Pvt Ltd", level=1)
    doc.add_heading("Vendor Contracts Digest and Approved Vendor List (FY 2026-27)", level=2)
    doc.add_heading("Clause VC-1 - Force Majeure (all Master Supply Agreements)", level=3)
    doc.add_paragraph("VC-1.1 A supplier affected by a force majeure event (natural disaster, port closure, war, blockade, government action) must notify Aravalli in writing within 48 hours.")
    doc.add_paragraph("VC-1.2 During force majeure Aravalli may source affected volumes from alternate suppliers without penalty and without breaching any volume commitment.")
    doc.add_paragraph("VC-1.3 If force majeure exceeds 30 days, Aravalli may terminate affected POs without liability.")
    doc.add_heading("Clause VC-2 - Late Delivery Penalty", level=3)
    doc.add_paragraph("VC-2.1 For delays not covered by force majeure, the supplier pays liquidated damages of 0.5% of the delayed PO value per week, capped at 5%.")
    doc.add_heading("Clause VC-3 - Primary Supplier Terms", level=3)
    doc.add_paragraph("VC-3.1 Formosa Microchip Co. (S101), Taiwan: primary for P-1001 and P-1002. Standard lead time 42 days ex-Kaohsiung. Supplier will re-route via Keelung on request at buyer cost, adding 4 days.")
    doc.add_paragraph("VC-3.2 Keelung Precision PCB Ltd. (S102), Taiwan: primary for P-2001. Lead time 35 days.")
    doc.add_paragraph("VC-3.3 Hsinchu SensorTech Inc. (S103), Taiwan: primary for P-3001 and P-3002. Ships via Kaohsiung.")
    doc.add_heading("Clause VC-4 - Backup Supplier Terms", level=3)
    doc.add_paragraph("VC-4.1 Penang Semicon Sdn. Bhd. (S201), Malaysia: pre-qualified backup for P-1001 and P-1002. Price premium up to 8% over primary. Standard lead time 21 days; expedited 10 days at +12% surcharge. Must acknowledge RFQs within 24 h.")
    doc.add_paragraph("VC-4.2 Saigon Circuit Works JSC (S202), Vietnam: backup for P-2001. Lead time 24 days; expedite 14 days at +10%.")
    doc.add_paragraph("VC-4.3 Bengaluru Embedded Systems (S301), India: domestic backup for P-1001, P-3001, P-3002. Lead time 14 days; capacity-limited (see ERP monthly_capacity). Preferred for L3 events because it removes sea-freight exposure.")
    doc.add_paragraph("VC-4.4 Pune CellWorks Pvt. Ltd. (S402), India: CONDITIONAL backup for P-5001 pending PPAP completion; any RFQ requires Head of SCM sign-off (see SOP-SC-014 clause 3.4).")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Schedule A - Approved Vendor List", level=3)
    add_table(doc,
              ["ID", "Supplier", "Country", "Status", "Parts covered"],
              [["S101", "Formosa Microchip Co.", "Taiwan", "APPROVED", "P-1001, P-1002 (primary)"],
               ["S102", "Keelung Precision PCB Ltd.", "Taiwan", "APPROVED", "P-2001 (primary)"],
               ["S103", "Hsinchu SensorTech Inc.", "Taiwan", "APPROVED", "P-3001, P-3002 (primary)"],
               ["S201", "Penang Semicon Sdn. Bhd.", "Malaysia", "APPROVED", "P-1001, P-1002 (backup)"],
               ["S202", "Saigon Circuit Works JSC", "Vietnam", "APPROVED", "P-2001 (backup)"],
               ["S301", "Bengaluru Embedded Systems", "India", "APPROVED", "P-1001, P-3001, P-3002 (backup)"],
               ["S302", "Chennai Interconnect Ltd.", "India", "APPROVED", "P-4001 (primary)"],
               ["S401", "Shenzhen PowerCell Co.", "China", "APPROVED", "P-5001 (primary)"],
               ["S402", "Pune CellWorks Pvt. Ltd.", "India", "CONDITIONAL", "P-5001 (backup)"]])
    doc.add_heading("Schedule B - RFQ Standard Terms", level=3)
    doc.add_paragraph("B.1 Quote validity 15 days.")
    doc.add_paragraph("B.2 Incoterms: CIF Nhava Sheva for imports, DAP Manesar for domestic.")
    doc.add_paragraph("B.3 Payment: 60 days from GRN.")
    doc.add_paragraph("B.4 Quotes must state lead time, MOQ and capacity commitment.")
    doc.save(path)
    print("wrote", path)


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    build_sop(os.path.join(here, "SOP-SC-014_Supply_Disruption_Response.docx"))
    build_vendor(os.path.join(here, "Vendor_Contracts_and_AVL.docx"))


if __name__ == "__main__":
    main()
